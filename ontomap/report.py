"""報告書（report/report.md と report.docx）。計画書 10 章の判定基準と 13 章の禁止事項を含める。"""
from __future__ import annotations

import json
import math
from pathlib import Path

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt  # noqa: E402
import numpy as np  # noqa: E402
import pandas as pd  # noqa: E402

plt.rcParams["font.family"] = ["Hiragino Sans", "Hiragino Kaku Gothic ProN", "sans-serif"]
plt.rcParams["axes.unicode_minus"] = False

SERIES = {"real": "#2a78d6", "paraphrase": "#1baf7a", "inverted": "#eb6834"}
TYPE_JA = {"real": "実文書", "paraphrase": "言い換え", "inverted": "関係反転"}
MARKERS = ["o", "s", "^", "D"]
FORBIDDEN = [
    "二次元マップで隣接することを、特許の侵害・新規性・同一技術性の証拠としない。",
    "合成文書の結果を実文書一般に外挿しない。",
    "正解グラフを用いた上限性能と、自動抽出を含む実運用性能を混同しない。",
    "2D マップの見かけのクラスター形状のみで品質を比較しない。",
    "各比較方式の前処理範囲・モデル読み込み条件を変えたまま「何倍速い」と記載しない。",
]


def _f(v, d=3) -> str:
    if v is None or (isinstance(v, float) and math.isnan(v)):
        return "—"
    return f"{v:.{d}f}"


def _md_table(df: pd.DataFrame) -> str:
    cols = list(df.columns)
    lines = ["| " + " | ".join(cols) + " |", "|" + "|".join("---" for _ in cols) + "|"]
    for r in df.itertuples(index=False):
        lines.append("| " + " | ".join(_f(x) if isinstance(x, float) else str(x) for x in r) + " |")
    return "\n".join(lines)


def _maps_figure(run, coords: pd.DataFrame, methods: list[str], path: Path) -> None:
    docs = {d.doc_id: d for d in run.docs}
    domains = sorted({d.tech_domain for d in run.docs})
    fig, axes = plt.subplots(1, len(methods), figsize=(4.2 * len(methods), 4.4))
    for ax, m in zip(np.atleast_1d(axes), methods):
        g = coords[(coords.method == m) & (coords.reducer == "mds")]
        for t, col in SERIES.items():
            for k, dom in enumerate(domains):
                sel = g[[docs[i].source_type == t and docs[i].tech_domain == dom for i in g.doc_id]]
                ax.scatter(sel.x, sel.y, s=22, c=col, marker=MARKERS[k % 4], edgecolors="white", linewidths=0.6, label=f"{TYPE_JA[t]}・{dom}")
        ax.set_title(m, fontsize=11)
        ax.set_xticks([]); ax.set_yticks([])
        for s in ax.spines.values():
            s.set_color("#e2e1dc")
    handles = [plt.Line2D([], [], marker="o", ls="", color=c, label=TYPE_JA[t]) for t, c in SERIES.items()]
    handles += [plt.Line2D([], [], marker=MARKERS[k % 4], ls="", color="#52514e", label=d) for k, d in enumerate(domains)]
    fig.legend(handles=handles, loc="lower center", ncol=len(handles), frameon=False, fontsize=9)
    fig.tight_layout(rect=(0, 0.08, 1, 1))
    fig.savefig(path, dpi=150)
    plt.close(fig)


def _movement_figure(run, coords: pd.DataFrame, a: str, b: str, path: Path) -> None:
    docs = {d.doc_id: d for d in run.docs}
    A = coords[(coords.method == a) & (coords.reducer == "mds")].set_index("doc_id")
    B = coords[(coords.method == b) & (coords.reducer == "mds")].set_index("doc_id").loc[A.index]
    fig, ax = plt.subplots(figsize=(6.4, 6.4))
    for i in A.index:
        ax.annotate("", xy=(B.x[i], B.y[i]), xytext=(A.x[i], A.y[i]), arrowprops=dict(arrowstyle="-|>", color="#9b9a94", lw=0.6, alpha=0.6, mutation_scale=6))
    for t, col in SERIES.items():
        sel = [i for i in A.index if docs[i].source_type == t]
        ax.scatter(B.x[sel], B.y[sel], s=20, c=col, edgecolors="white", linewidths=0.6, label=TYPE_JA[t], zorder=3)
    ax.set_title(f"{a} の位置 → {b} の位置（MDS、Procrustes で整列）", fontsize=10)
    ax.set_xticks([]); ax.set_yticks([])
    ax.legend(frameon=False, fontsize=9, loc="lower right")
    fig.tight_layout()
    fig.savefig(path, dpi=150)
    plt.close(fig)


def _criteria(q: pd.DataFrame, fid: pd.DataFrame, bench: pd.DataFrame | None, gq: dict, primary: str, ref: str = "B2") -> list[tuple[str, str, str]]:
    """計画書 10 章の判定。返り値: (判定項目, 結果, 根拠)。"""
    out = []
    qa = q[q.subset == "all"].set_index("method")
    if primary in qa.index and ref in qa.index:
        p, r = qa.loc[primary], qa.loc[ref]
        d_rho = p.spearman - r.spearman
        d_ndcg = p["ndcg@10"] - r["ndcg@10"]
        ok = d_rho >= 0.08 or d_ndcg >= 0.05
        ci = f"Spearman 差の 95% 区間 [{_f(p.get(f'spearman_minus_{ref}_ci_low'))}, {_f(p.get(f'spearman_minus_{ref}_ci_high'))}]、nDCG 差の 95% 区間 [{_f(p.get(f'ndcg_minus_{ref}_ci_low'))}, {_f(p.get(f'ndcg_minus_{ref}_ci_high'))}]"
        out.append(("品質上の成功（Spearman +0.08 以上 または nDCG@10 +0.05 以上）", "達成" if ok else "未達", f"{primary} − {ref}: Spearman {_f(d_rho)}、nDCG@10 {_f(d_ndcg)}。{ci}"))
        layer = primary.split("@")[1] if "@" in primary else ""
        a1, a2 = qa.loc[f"A1@{layer}"] if f"A1@{layer}" in qa.index else None, qa.loc[f"A2@{layer}"] if f"A2@{layer}" in qa.index else None
        if a1 is not None and a2 is not None:
            beats = p["ndcg@10"] > a1["ndcg@10"] and p["ndcg@10"] > a2["ndcg@10"]
            inv_low = p.get("inv_percentile", np.nan) < min(a1.get("inv_percentile", np.nan), r.get("inv_percentile", np.nan))
            out.append(("関係性の有効性（P2 が A1/A2 を上回り、関係違いの負例を低く順位づける）", "達成" if beats and inv_low else "一部未達" if beats or inv_low else "未達",
                        f"nDCG@10: P2 {_f(p['ndcg@10'])} / A1 {_f(a1['ndcg@10'])} / A2 {_f(a2['ndcg@10'])}。関係反転文書の類似順位の百分位（低いほど良い）: P2 {_f(p.get('inv_percentile'))} / A1 {_f(a1.get('inv_percentile'))} / {ref} {_f(r.get('inv_percentile'))}"))
    fm = fid[fid.reducer == "mds"].set_index("method")
    if primary in fm.index and ref in fm.index:
        d = fm.loc[primary, "knn@10"] - fm.loc[ref, "knn@10"]
        out.append(("二次元配置（kNN overlap@10 が B2 より明確に改善、または同程度で説明可能性が向上）", "達成" if d >= 0.05 else "同程度" if abs(d) < 0.05 else "未達", f"MDS の kNN overlap@10: {primary} {_f(fm.loc[primary, 'knn@10'])} / {ref} {_f(fm.loc[ref, 'knn@10'])}（差 {_f(d)}）。説明可能性は人手評価が必要"))
    if bench is not None and not bench.empty:
        e2e = bench[(bench.stage == "end_to_end_mds") & (bench.method == "P2")]
        if not e2e.empty:
            v = float(e2e["median"].iloc[0])
            out.append(("運用可能性（辞書抽出ベースで End-to-end 10 秒以内）", "達成" if v <= 10 else "未達", f"P2（辞書抽出→類似度→MDS）の warm 中央値 {_f(v, 2)} 秒。環境依存の目標値"))
    if gq:
        layer = next(iter(gq))
        out.append(("不成功の解釈（抽出がボトルネックか）", "参考", f"{layer} の三つ組 F1 {_f(gq[layer]['triple_f1'])}、概念 F1 {_f(gq[layer]['concept_f1'])}。正解グラフ（gold）との差が抽出の影響"))
    return out


def build_report(run) -> Path:
    """report/report.md と report/report.docx を書く。"""
    from . import pipeline as pl

    pl.ensure_docs(run)
    rd = run.root / "report"
    rd.mkdir(exist_ok=True)
    q = pd.read_csv(run.root / "eval" / "similarity_quality.csv")
    fid = pd.read_csv(run.root / "eval" / "map_fidelity.csv")
    lp = pd.read_csv(run.root / "eval" / "map_label_preservation.csv")
    gq = json.loads((run.root / "eval" / "graph_quality.json").read_text(encoding="utf-8"))
    tw = json.loads((run.root / "eval" / "tuned_weights.json").read_text(encoding="utf-8")) if (run.root / "eval" / "tuned_weights.json").exists() else {}
    q300 = pd.read_csv(run.root / "eval" / "similarity_quality_300pairs.csv") if (run.root / "eval" / "similarity_quality_300pairs.csv").exists() else None
    mv = pd.read_csv(run.root / "maps" / "movement.csv") if (run.root / "maps" / "movement.csv").exists() else pd.DataFrame()
    bench = pd.read_csv(run.root / "bench" / "benchmark_summary.csv") if (run.root / "bench" / "benchmark_summary.csv").exists() else None
    benv = json.loads((run.root / "bench" / "environment.json").read_text(encoding="utf-8")) if (run.root / "bench" / "environment.json").exists() else {}
    man = json.loads((run.root / "manifest.json").read_text(encoding="utf-8")) if (run.root / "manifest.json").exists() else {}
    coords = pd.read_csv(run.root / "maps" / "map_coords.csv")
    synthetic = (run.root / "data" / "_truth.jsonl").exists()
    label_src = q.label_source.iloc[0]
    primary = "P2@dict" if "P2@dict" in set(q.method) else "P2@gold"

    fig_methods = [m for m in ["B1", "B2", "P2@dict", "P3@dict"] if m in set(coords.method)]
    _maps_figure(run, coords, fig_methods, rd / "fig_maps_mds.png")
    if "B2" in set(coords.method) and primary in set(coords.method):
        _movement_figure(run, coords, "B2", primary, rd / "fig_movement.png")

    crit = _criteria(q, fid, bench, gq, primary)
    order = ["B1", "B2", "B3", "P1@gold", "P2@gold", "P3@gold", "P1@dict", "P2@dict", "P3@dict", "P3tuned@dict", "A1@dict", "A2@dict", "A1@gold", "A2@gold", "P3tuned@gold"]
    order += [m for m in q.method.unique() if m not in order]

    def qtab(subset: str, cols: list[str]) -> pd.DataFrame:
        t = q[q.subset == subset].set_index("method")
        t = t.loc[[m for m in order if m in t.index], cols].reset_index()
        return t

    qa = qtab("all", ["coverage", "spearman", "spearman_ci_low", "spearman_ci_high", "ndcg@10", "recall@10"])
    qa.columns = ["手法", "比較可能率", "Spearman", "95%下限", "95%上限", "nDCG@10", "Recall@10"]
    qr = qtab("real_only", ["spearman", "ndcg@10", "recall@10"])
    qr.columns = ["手法", "Spearman", "nDCG@10", "Recall@10"]
    ctl = qtab("all", ["para_mrr", "para_hit10", "inv_percentile", "para_beats_inv"]) if "para_mrr" in q else None
    if ctl is not None:
        ctl.columns = ["手法", "言い換え MRR", "言い換え Hit@10", "反転文書の百分位（低いほど良い）", "言い換え＞反転の率"]
    fm = fid.copy()
    fm = fm.merge(lp[["method", "reducer", "relevant_in_2d_knn10"]], on=["method", "reducer"], how="left")
    fm = fm.set_index(["method"]).loc[[m for m in order if m in set(fm.method)]].reset_index()
    fmt = fm[["method", "reducer", "knn@5", "knn@10", "knn@20", "spearman_dist", "stress", "relevant_in_2d_knn10"]]
    fmt.columns = ["手法", "配置", "kNN@5", "kNN@10", "kNN@20", "距離 Spearman", "ストレス", "スコア3ペアの2D近傍残存率"]

    L: list[str] = []
    L.append(f"# オントロジー型文書類似度と二次元マッピングの比較実験 報告書（{run.root.name}）\n")
    if synthetic:
        L.append("> **注意**: この結果は見本コーパス（合成文書 200 件）と、生成時の構造から機械的に決めた模擬ラベルによるものです。"
                 "パイプラインの動作確認であり、仮説 H1〜H3 の検証結果ではありません。模擬ラベルと正解グラフは同じ生成構造に由来するため、"
                 "関係方式に有利な循環があります。辞書抽出の辞書と生成語彙も同じオントロジーなので、概念 F1 は 1.0 になります。\n")
    L.append(f"ラベル: **{label_src}**（human = 専門家の採点、simulated = 模擬ラベル）\n")
    L.append("## 1. 要点\n")
    for name, res, why in crit:
        L.append(f"- **{name}**: {res}。{why}")
    if not mv.empty and primary in set(mv["to"]):
        m = mv[(mv["from"] == "B2") & (mv["to"] == primary) & (mv.reducer == "mds")]
        L.append(f"- **B2 から {primary} への配置の変化**: 2D の 10 近傍の重なりは平均 {_f(m.knn_jaccard.mean())}（1 = 同じ近傍）。移動距離の中央値は {_f(m['shift'].median())}（配置半径 = 1 の単位）。")
    L.append("")
    L.append("## 2. データ\n")
    st = pd.Series([d.source_type for d in run.docs]).value_counts()
    ln = pd.Series([d.char_count for d in run.docs])
    L.append(f"- 文書数 {len(run.docs)}（{'、'.join(f'{TYPE_JA.get(k, k)} {v}' for k, v in st.items())}）。ペア数 {len(run.docs) * (len(run.docs) - 1) // 2}。")
    L.append(f"- 文字数 最小 {ln.min()}・中央値 {int(ln.median())}・最大 {ln.max()}。")
    L.append("- 重み調整は train 群の文書どうしのペアだけを使い、評価は残りのペア（eval）で行った。文書群は見本ではメカニズム単位。\n")
    L.append("## 3. グラフ抽出の品質（正解グラフ基準）\n")
    if gq:
        g = pd.DataFrame(gq).T.reset_index().rename(columns={"index": "層"})
        L.append(_md_table(g[["層", "concept_f1", "triple_precision", "triple_recall", "triple_f1", "triple_f1_with_role_negation"]]))
    L.append("\n## 4. 類似度の品質（eval ペア）\n")
    L.append("### 4.1 全文書\n")
    L.append(_md_table(qa))
    L.append("\n### 4.2 実文書だけ\n")
    L.append(_md_table(qr))
    if ctl is not None:
        L.append("\n### 4.3 制御文書（H1: 言い換え、H2: 関係反転）\n")
        L.append(_md_table(ctl))
    if q300 is not None:
        L.append("\n### 4.4 層化抽出 300 ペアでの評価（採点済み候補だけで順位を作る）\n")
        L.append(_md_table(q300.set_index("method").loc[[m for m in order if m in set(q300.method)]].reset_index().rename(columns={"method": "手法"})))
    if tw:
        L.append(f"\n重み調整（train ペア {tw.get('n_train_pairs')} 件、Spearman 最大化）の結果: " + "、".join(f"{k}: α,β,γ = {', '.join(_f(x, 1) for x in v['weights'])}" for k, v in tw.items() if isinstance(v, dict)))
    L.append("\n## 5. 二次元マップの忠実度\n")
    L.append(_md_table(fmt))
    L.append("\n![MDS マップ](fig_maps_mds.png)\n")
    if (rd / "fig_movement.png").exists():
        L.append(f"![移動](fig_movement.png)\n")
    if not mv.empty:
        L.append("### 手法間の移動（MDS、B2 から）\n")
        mm = mv[(mv["from"] == "B2") & (mv.reducer == "mds")].groupby("to")[["shift", "knn_jaccard"]].mean().reset_index()
        mm.columns = ["移動先の手法", "平均移動距離", "2D 10 近傍の重なり"]
        L.append(_md_table(mm))
    if bench is not None:
        L.append("\n## 6. 計算時間とメモリ\n")
        L.append(f"環境: {benv.get('cpu', '')}、メモリ {benv.get('memory_gb', '')} GB、スレッド {benv.get('threads', '')}、{benv.get('device', '')}。warm は warm-up 後の本計測、cold は新しいプロセスでの初回（モデル読込・JIT を含む）。\n")
        b = bench.copy()
        b = b[b.stage.isin(["import", "model_load", "features", "pairwise", "mds", "umap", "end_to_end_mds", "extract"])]
        b = b[["method", "stage", "median", "p2_5", "p97_5", "n", "cold_seconds", "peak_ram_mb"]]
        b["n"] = b["n"].map(lambda v: "—" if pd.isna(v) else str(int(v)))
        b["peak_ram_mb"] = b["peak_ram_mb"].map(lambda v: "—" if pd.isna(v) else f"{v:.0f}")
        b.columns = ["手法", "工程", "warm 中央値(秒)", "2.5%", "97.5%", "回数", "cold(秒)", "プロセスのピーク RSS(MB)"]
        L.append(_md_table(b))
        L.append("\nピーク RSS はプロセス全体の最大値で、工程ごとの値ではない。")
    L.append("\n## 7. 結果を解釈するときの禁止事項\n")
    for x in FORBIDDEN:
        L.append(f"- {x}")
    L.append("\n## 8. 再現情報\n")
    if man:
        L.append(f"- 設定の SHA-256: `{man.get('config_sha256', '')}`")
        L.append(f"- Python {man.get('python')}、{man.get('platform')}")
        L.append("- ライブラリ: " + "、".join(f"{k} {v}" for k, v in man.get("libraries", {}).items() if v))
    sm = json.loads((run.root / "similarity" / "manifest.json").read_text(encoding="utf-8"))
    if "B2" in sm:
        L.append(f"- 文埋め込み: {sm['B2'].get('model')}（revision {sm['B2'].get('revision')}、{sm['B2'].get('dim')} 次元）")
    L.append(f"- 乱数種: {run.config['seed']}")
    md = "\n".join(L) + "\n"
    (rd / "report.md").write_text(md, encoding="utf-8")
    _write_docx(md, rd)
    return rd / "report.md"


def _write_docx(md: str, rd: Path) -> None:
    """report.md の内容を docx に写す（見出し・箇条書き・表・図）。"""
    from docx import Document
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.oxml.ns import qn
    from docx.shared import Inches, Pt

    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = "Hiragino Sans"
    st.element.rPr.rFonts.set(qn("w:eastAsia"), "Hiragino Sans")
    st.font.size = Pt(10)
    lines = md.splitlines()
    i = 0
    while i < len(lines):
        s = lines[i]
        if s.startswith("# "):
            doc.add_heading(s[2:], level=0)
        elif s.startswith("## "):
            doc.add_heading(s[3:], level=1)
        elif s.startswith("### "):
            doc.add_heading(s[4:], level=2)
        elif s.startswith("|"):
            rows = []
            while i < len(lines) and lines[i].startswith("|"):
                if not set(lines[i].replace("|", "").strip()) <= set("-"):
                    rows.append([c.strip() for c in lines[i].strip("|").split("|")])
                i += 1
            t = doc.add_table(rows=len(rows), cols=len(rows[0]))
            t.style = "Light Grid Accent 1"
            t.alignment = WD_TABLE_ALIGNMENT.CENTER
            for r, row in enumerate(rows):
                for c, v in enumerate(row):
                    cell = t.cell(r, c)
                    cell.text = v
                    for p in cell.paragraphs:
                        for run_ in p.runs:
                            run_.font.size = Pt(8)
            continue
        elif s.startswith("![") and "](" in s:
            path = rd / s[s.index("](") + 2 : s.rindex(")")]
            if path.exists():
                doc.add_picture(str(path), width=Inches(6.3))
        elif s.startswith("- "):
            doc.add_paragraph(s[2:].replace("**", ""), style="List Bullet")
        elif s.startswith("> "):
            p = doc.add_paragraph(s[2:].replace("**", ""))
            p.runs[0].italic = True
        elif s.strip():
            doc.add_paragraph(s.replace("**", "").replace("`", ""))
        i += 1
    doc.save(rd / "report.docx")
